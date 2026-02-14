from PyQt5.QtWidgets import QMessageBox
from docx import Document
import spacy
from indexes import gulpease_index, lexical_diversity_index, syntactic_complexity_index
from langdetect import detect, DetectorFactory

nlp = spacy.load('it_core_news_sm')


def extract_t_units(sent):
    t_units = []
    roots = [token for token in sent if token.dep_ == "ROOT"]
    for root in roots:
        # subtree del root + coordinate
        tokens = list(root.subtree)
        for child in root.children:
            if child.dep_ == "conj":  # coordinate
                tokens.extend(list(child.subtree))
        tokens = sorted(set(tokens), key=lambda t: t.i)
        t_units.append(tokens)
    return t_units


class WordFile:
    def __init__(self, file_path=None, text=""):
        self.file_path = file_path if file_path else None
        self.doc = Document(file_path) if file_path else None
        if file_path:
            self.paragraphs = self.doc.paragraphs
            self.plain_text = ""
            for para in self.paragraphs:
                self.plain_text += para.text + "\n"
        else:
            self.plain_text = text
        self.tokens = nlp(self.plain_text)
        self.letters_num = sum(len(token.text) for token in self.tokens if token.is_alpha)
        self.words = [token.text.lower() for token in self.tokens if token.is_alpha]
        self.unique_words_num = len(set(self.words))
        self.words_num = len(self.words)
        self.sentences = list(self.tokens.sents)
        self.sentences_num = len(list(self.tokens.sents))

        self.subclauses_num = 0
        self.mainclauses_num = 0
        all_t_units = []
        for i, sent in enumerate(self.sentences, 1):
            t_units = extract_t_units(sent)
            for tokens in t_units:
                t_unit_text = " ".join([t.text for t in tokens])
                all_t_units.append(t_unit_text)

                root_verbs = [t for t in tokens if t.dep_ == "ROOT"]
                sub_clauses = [t for t in tokens if t.dep_ in ("advcl", "ccomp", "acl", "xcomp")]
                self.subclauses_num += len(sub_clauses)
                self.mainclauses_num += len(root_verbs)
        self.clauses_num = self.mainclauses_num + self.subclauses_num
        self.t_units_num = len(all_t_units)
        if len(self.plain_text) > 200:
            self.lang = detect(self.plain_text)
        else:
            self.lang = None


    def calculate_gix(self):
        return gulpease_index.calculate(self.sentences_num, self.letters_num, self.words_num)

    def calculate_scix(self):
        return syntactic_complexity_index.calculate(self.sentences_num, self.words_num, self.clauses_num, self.mainclauses_num, self.t_units_num)

    def calculate_lix(self):
        return lexical_diversity_index.calculate(self.unique_words_num, self.words_num)